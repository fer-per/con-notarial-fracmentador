"""
Repositorio Excel — implementación del ExcelRepositoryPort.

Carga inventarios archivísticos desde archivos Excel (.xlsx)
usando pandas + openpyxl. El mapeo de columnas, metadatos, hoja y
reglas de filtrado provienen del modelo de catálogo activo.
"""
import logging
import re
import unicodedata
from dataclasses import fields
from datetime import date, datetime
from typing import List, Optional

import pandas as pd
from openpyxl import load_workbook

from src.domain.entities import InventoryRecord
from src.domain.modelo_catalogo import ModeloCatalogo, UbicacionMetadato
from src.domain.ports.excel_port import ExcelRepositoryPort
from src.domain.value_objects import (
    SKIP_ROW_KEYWORDS,
    ANNOTATION_KEYWORDS,
    EMPTY_VALUES,
)
from src.infrastructure.modelo_repository import ModeloRepository

logger = logging.getLogger(__name__)

_CAMPOS_INVENTARIO = {campo.name for campo in fields(InventoryRecord)}


class ExcelRepository(ExcelRepositoryPort):
    """Implementación concreta del repositorio Excel."""

    def __init__(self, modelo: Optional[ModeloCatalogo] = None):
        self._modelo = modelo or ModeloRepository().cargar("historico")

    def set_modelo(self, modelo: ModeloCatalogo):
        """Cambia el modelo de catálogo usado por este repositorio."""
        self._modelo = modelo

    def cargar_registros(
        self,
        ruta: str,
        fila_datos_inicio: int,
        fila_inicio: int,
        fila_fin: int,
    ) -> List[InventoryRecord]:
        """Carga registros del inventario Excel según el modelo activo."""
        logger.info("Cargando registros desde: %s (modelo %s)", ruta, self._modelo.id)

        excel = self._modelo.excel
        filas_encabezado = excel.filas_encabezado
        skiprows = fila_datos_inicio - filas_encabezado - 1

        df = pd.read_excel(
            ruta,
            sheet_name=excel.hoja,
            skiprows=skiprows,
            header=list(range(filas_encabezado)),
            engine='openpyxl',
        )

        # Aplanar MultiIndex de columnas
        df.columns = [
            ' '.join(str(c) for c in col).strip()
            if isinstance(col, tuple) else str(col)
            for col in df.columns
        ]

        # Mapear columnas
        col_map = self._mapear_columnas(df.columns.tolist())
        logger.info("Columnas mapeadas: %s", col_map)

        records: List[InventoryRecord] = []
        record_counter = 0

        anotaciones = excel.anotaciones or tuple(ANNOTATION_KEYWORDS)
        prefijos = excel.ignorar_prefijos or tuple(SKIP_ROW_KEYWORDS)
        regex_ignorar = [
            re.compile(patron, re.IGNORECASE)
            for patron in excel.ignorar_regex
        ]

        # Filas de anotación/instrucciones (no son datos)
        annot = df.apply(
            lambda row: any(
                kw in self._safe_str(v).lower()
                for v in row for kw in anotaciones
            ),
            axis=1,
        ).tolist()

        rows = list(df.iterrows())
        for pos, (_, row) in enumerate(rows):
            fila_real = fila_datos_inicio + pos

            # Filtrar por rango de filas
            if fila_real < fila_inicio or fila_real > fila_fin:
                continue

            # Verificar si es fila de anotación/instrucción
            if annot[pos]:
                continue

            # Verificar si es fila de ejemplo (seguida de una anotación)
            if pos + 1 < len(rows) and annot[pos + 1]:
                continue

            # Verificar sub-encabezados y patrones configurados
            first_cell = self._safe_str(row.iloc[0]) if len(row) else ""
            if any(kw in first_cell.lower() for kw in prefijos):
                continue
            if any(rx.search(first_cell) for rx in regex_ignorar):
                continue

            # Extraer todos los campos declarados por el modelo
            valores = {
                campo: self._get_col(row, col_map, campo)
                for campo in self._modelo.columnas
                if campo in _CAMPOS_INVENTARIO
            }

            # Filtrar filas fantasma: sin ningún campo fuerte
            if excel.ignorar_si_sin_fuertes:
                fuertes = [
                    valores.get(campo, "")
                    for campo in excel.campos_fuertes
                ]
                if not any(valor.strip() for valor in fuertes):
                    continue

            record_counter += 1
            for campo in ("registro", "escribano", "folios", "titulo"):
                valores.setdefault(campo, "")
            valores["protocolo"] = self._normalizar_protocolo(
                valores.get("protocolo", "")
            )
            for campo in ("tomo", "escritura"):
                if campo in valores:
                    valores[campo] = self._normalizar_protocolo(valores[campo])

            records.append(InventoryRecord(
                id=f"#{record_counter:04d}",
                fila=fila_real,
                pg_pdf="",  # Se calcula después con el mapper
                **valores,
            ))

        logger.info("Cargados %d registros.", len(records))
        return records

    def extraer_metadatos(self, ruta: str, fila_datos_inicio: int) -> dict:
        """Extrae los metadatos globales declarados por el modelo activo."""
        logger.info("Extrayendo metadatos de: %s (modelo %s)", ruta, self._modelo.id)

        metadatos = {
            "filepath": ruta,
            "siglo": "",
            "escribano": "",
            "acervo_num": "",
        }

        try:
            excel = self._modelo.excel
            filas_encabezado = excel.filas_encabezado
            nrows = max(0, fila_datos_inicio - filas_encabezado - 1)
            df_header = pd.read_excel(
                ruta,
                sheet_name=excel.hoja,
                header=None,
                nrows=nrows,
                engine='openpyxl',
            )

            for campo, ubicacion in self._modelo.metadatos.items():
                cell = self._celda(df_header, ubicacion.fila, ubicacion.columna)
                metadatos[campo] = self._aplicar_transformacion(cell, ubicacion)

        except Exception as e:
            logger.warning("Error extrayendo metadatos: %s", e)

        # Aliases de compatibilidad usados por los casos de uso actuales.
        metadatos["acervo_num"] = (
            metadatos.get("acervo", "") or metadatos.get("acervo_num", "")
        )
        metadatos["siglo"] = metadatos.get("siglo", "")
        metadatos["escribano"] = metadatos.get("escribano", "")
        return metadatos

    @staticmethod
    def _celda(df_header, fila: int, columna: int) -> str:
        """Lee una celda 1-based del DataFrame de cabecera como string."""
        if fila - 1 >= len(df_header):
            return ""
        row = df_header.iloc[fila - 1]
        if columna - 1 >= len(row):
            return ""
        return ExcelRepository._safe_str(row.iloc[columna - 1])

    @staticmethod
    def _aplicar_transformacion(cell: str, ubicacion: UbicacionMetadato) -> str:
        """Aplica regex y transformación declaradas para un metadato."""
        valor = (cell or "").strip()
        if ubicacion.regex:
            match = re.search(ubicacion.regex, valor, re.IGNORECASE)
            if not match:
                return ""
            try:
                valor = match.group(ubicacion.grupo) or ""
            except (IndexError, TypeError):
                valor = match.group(0)

        transformacion = ubicacion.transformacion or "texto"
        if transformacion == "texto":
            valor = ExcelRepository._extraer_texto(valor)
        elif transformacion == "numero":
            match = re.search(r"\d+", valor)
            valor = match.group(0) if match else ""
        elif transformacion == "siglo_romano":
            valor = valor.strip().upper()
            if not re.fullmatch(r"[IVXLCDM]+", valor):
                return ""
        elif transformacion == "texto_completo":
            valor = valor.strip()

        if ubicacion.prefijo:
            valor = f"{ubicacion.prefijo}{valor}"
        return valor.strip()

    @staticmethod
    def _extraer_texto(cell: str) -> str:
        """Obtiene el texto posterior a ':' / '—' / '-' o el valor completo."""
        if not cell:
            return ""
        for sep in (":", "—", "-"):
            if sep in cell:
                return cell.split(sep, 1)[1].strip()
        return cell.strip()

    def detectar_fila_inicio_datos(self, ruta: str) -> Optional[int]:
        """Detecta la fila 1-based donde empiezan los datos según el modelo."""
        excel = self._modelo.excel
        keywords = excel.detectar_keywords or (
            "registro", "escribano", "protocolo", "folios",
        )
        minimo = max(1, excel.detectar_min_matches)

        try:
            df = pd.read_excel(
                ruta,
                sheet_name=excel.hoja,
                header=None,
                nrows=200,
                engine='openpyxl',
            )
        except Exception as e:
            logger.warning("No se pudo detectar la fila de datos: %s", e)
            return None

        for idx, row in df.iterrows():
            cells = [
                self._sin_acentos(self._safe_str(v).lower())
                for v in row if v is not None
            ]
            non_empty = [c for c in cells if c]
            if len(non_empty) < 3:
                continue
            matches = sum(
                1 for kw in keywords
                for c in non_empty
                if self._sin_acentos(kw.lower()) in c
            )
            if matches >= minimo:
                return int(idx) + 1 + excel.filas_encabezado

        return None

    def guardar_registros(
        self,
        ruta: str,
        fila_datos_inicio: int,
        records: List[InventoryRecord],
    ) -> int:
        """Escribe de vuelta los valores de los registros en el Excel.

        Solo se sobrescriben las columnas localizadas por el mapa maestro
        y los registros con ``fila`` válida. Las celdas cuyo valor en el
        registro está vacío no se tocan, para no borrar contenido que no
        fue parseado.
        """
        try:
            wb = load_workbook(ruta)
        except Exception as e:
            logger.error("No se pudo abrir el Excel para guardar: %s", e)
            raise

        hoja = self._modelo.excel.hoja
        if isinstance(hoja, str):
            ws = wb[hoja]
        elif isinstance(hoja, int):
            ws = wb.worksheets[hoja]
        else:
            ws = wb.worksheets[0]
        columnas = self._indices_columnas(ws, fila_datos_inicio)

        celdas = 0
        for record in records:
            fila = record.fila
            if not fila:
                continue
            for campo, col in columnas.items():
                valor = getattr(record, campo, "")
                if valor in (None, ""):
                    continue
                ws.cell(row=fila, column=col).value = valor
                celdas += 1

        wb.save(ruta)
        logger.info("Guardados %d valores en el Excel %s.", celdas, ruta)
        return celdas

    def _indices_columnas(self, ws, fila_datos_inicio: int) -> dict:
        """Localiza la columna (1-based) de cada campo en las filas de encabezado.

        Recrea el texto combinado de cada columna a partir de las filas
        de encabezado declaradas por el modelo y delega en
        ``_mapear_columnas`` para usar exactamente los mismos aliases.
        """
        ncols = ws.max_column or 0
        filas_encabezado = self._modelo.excel.filas_encabezado
        encabezado = []
        for c in range(1, ncols + 1):
            partes = [
                self._safe_str(ws.cell(row=r, column=c).value)
                for r in range(
                    fila_datos_inicio - filas_encabezado,
                    fila_datos_inicio,
                )
            ]
            encabezado.append(" ".join(partes).strip())

        # Reutiliza el mismo mapeo que la lectura para que guardar y cargar
        # coincidan incluso cuando hay aliases cortos o solapados.
        indices = self._mapear_columnas(encabezado)
        return {campo: idx + 1 for campo, idx in indices.items()}

    def _mapear_columnas(self, columnas: List[str]) -> dict:
        """Mapea nombres de columnas del Excel a campos del modelo.

        Usa el mapeo maestro (aliases + índice de respaldo). Si no hay un
        campo mapeado, se lo omite del resultado.
        """
        col_map = {}
        columnas_lower = [self._sin_acentos(c.lower()) for c in columnas]

        for campo, mapeo in self._modelo.columnas.items():
            for alias in mapeo.aliases:
                alias_norm = self._sin_acentos(alias)
                for i, col_lower in enumerate(columnas_lower):
                    if alias_norm in col_lower:
                        col_map[campo] = i
                        break
                if campo in col_map:
                    break

            # Fallback por índice
            if campo not in col_map and mapeo.indice_fallback is not None:
                if mapeo.indice_fallback < len(columnas):
                    col_map[campo] = mapeo.indice_fallback

        return col_map

    @staticmethod
    def _sin_acentos(texto: str) -> str:
        """Elimina acentos/diacríticos para comparación insensible a acentos."""
        return ''.join(
            c for c in unicodedata.normalize('NFD', texto)
            if unicodedata.category(c) != 'Mn'
        )

    def _get_col(self, row, col_map: dict, campo: str) -> str:
        """Obtiene el valor de una columna mapeada como string limpio."""
        if campo not in col_map:
            return ""
        idx = col_map[campo]
        if idx >= len(row):
            return ""
        val = row.iloc[idx]
        return self._safe_str(val)

    @staticmethod
    def _safe_str(val) -> str:
        """Convierte un valor a string, tratando NaN/None como vacío."""
        if val is None:
            return ""
        try:
            if pd.isna(val):
                return ""
        except (TypeError, ValueError):
            pass
        if isinstance(val, (datetime, date)):
            return val.strftime("%Y-%m-%d")
        s = str(val).strip()
        if s.lower() in EMPTY_VALUES:
            return ""
        return s

    @staticmethod
    def _normalizar_protocolo(val: str) -> str:
        """Normaliza protocolo: si es float entero (ej: '3.0'), convierte a '3'."""
        if not val:
            return ""
        try:
            f = float(val)
            if f == int(f):
                return str(int(f))
        except (ValueError, TypeError):
            pass
        return val
