"""Tests del modelo notarial: herencia, detección, lectura y jerarquía."""
import os
import shutil
import sys
import tempfile
from datetime import datetime

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))

from openpyxl import Workbook

from src.infrastructure.excel_repository import ExcelRepository
from src.infrastructure.hierarchy_builder import HierarchyBuilder
from src.infrastructure.modelo_repository import ModeloRepository
from src.domain.entities import InventoryRecord


def _crear_notarial(path: str):
    """Crea un catálogo notarial mínimo con metadatos, fantasmas y dos registros."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja 1"

    ws.cell(row=4, column=1, value="REPOSITORIO DOCUMENTAL: N° 1")
    ws.cell(row=5, column=1, value="Fondo: Javier de Taboada Vizcarra")
    ws.cell(row=6, column=1, value="Entidad Productora: Notaria Javier de Taboada Vizcarra")
    ws.cell(row=7, column=1, value="Sección: Escrituras Públicas")
    ws.cell(row=8, column=1, value="Serie: Protocolos Notariales")
    ws.cell(row=9, column=1, value="Productor: Javier de Taboada Vizcarra")

    headers = [
        "N° DE PROT.", "N° DE TOMO", "N° DE ESCRITURA", "N° DE FOLIOS",
        "TIPO DE ESCRITURA", "OTORGANTE/INTERESADO 1",
        "OTORGANTE/INTERESADO 2", "OTORGANTE/INTERESADO 3",
        "FECHA INICIAL", "FECHA FINAL", "DATA TOPICA",
        "ALCANCE Y CONTENIDO", "ESTADO DE CONSERVACIÓN",
        "FECHA DE LA DESCRIPCIÓN", "NOTARIO", "OBSERVACIONES",
    ]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=11, column=col, value=header)
        ws.cell(row=12, column=col, value="N°")

    datos = [
        {
            "protocolo": 1, "tomo": 1, "escritura": 1, "folios": "1r",
            "titulo": "Protesto", "interesado1": "Oscar Gutierrez",
            "interesado2": "Adelina Tejada Jara", "interesado3": "",
            "fecha_inicio": datetime(1967, 12, 20),
            "fecha_fin": datetime(1967, 12, 20),
            "data_topica": "AREQUIPA",
            "alcance": "Protesto por falta de pago",
            "estado": "BUENO",
            "fecha_descripcion": datetime(2026, 3, 25),
            "escribano": "Javier de Taboada Vizcarra",
            "observaciones": "Sin observaciones",
        },
        {
            "protocolo": 1, "tomo": 1, "escritura": 2, "folios": "1v - 2r",
            "titulo": "Protesto", "interesado1": "Banco del Sur del Perú",
            "interesado2": "Casa Ynglesa S. A.", "interesado3": "Fábrica Moderna",
            "fecha_inicio": datetime(1967, 12, 20),
            "fecha_fin": datetime(1967, 12, 20),
            "data_topica": "AREQUIPA",
            "alcance": "Protesto por falta de pago",
            "estado": "BUENO",
            "fecha_descripcion": datetime(2026, 3, 25),
            "escribano": "Javier de Taboada Vizcarra",
            "observaciones": "",
        },
    ]
    for offset, dato in enumerate(datos):
        row = 16 + offset
        for col, key in enumerate([
            "protocolo", "tomo", "escritura", "folios", "titulo",
            "interesado1", "interesado2", "interesado3", "fecha_inicio",
            "fecha_fin", "data_topica", "alcance", "estado",
            "fecha_descripcion", "escribano", "observaciones",
        ], start=1):
            ws.cell(row=row, column=col, value=dato[key])

    # Fila fantasma: solo protocolo/tomo, sin campos fuertes.
    ws.cell(row=18, column=1, value=1)
    ws.cell(row=18, column=2, value=1)

    wb.save(path)


class TestModeloNotarial:
    def setup_method(self):
        self.repo = ModeloRepository()
        self._tmp_dir = tempfile.mkdtemp()
        self.tmp = os.path.join(self._tmp_dir, "notarial.xlsx")
        _crear_notarial(self.tmp)

    def teardown_method(self):
        shutil.rmtree(self._tmp_dir, ignore_errors=True)

    def test_hereda_columnas_base_y_declara_tomo(self):
        modelo = self.repo.cargar("notarial")
        assert modelo.id == "notarial"
        assert modelo.excel.hoja == "Hoja 1"
        assert "registro" in modelo.columnas  # heredado de base.json
        assert "tomo" in modelo.columnas
        assert len(modelo.jerarquia.niveles) == 12
        assert modelo.jerarquia.niveles[6].tipo == "tomo"

    def test_detecta_notarial_por_encabezados(self):
        assert self.repo.detectar(self.tmp) == "notarial"

    def test_extrae_metadatos_notariales(self):
        excel = ExcelRepository(self.repo.cargar("notarial"))
        fila = excel.detectar_fila_inicio_datos(self.tmp)
        metadatos = excel.extraer_metadatos(self.tmp, fila)

        assert fila == 13
        assert metadatos["acervo_num"] == "1"
        assert metadatos["fondo"] == "Javier de Taboada Vizcarra"
        assert metadatos["entidad"] == "Notaria Javier de Taboada Vizcarra"
        assert metadatos["seccion"] == "Escrituras Públicas"
        assert metadatos["serie"] == "Protocolos Notariales"
        assert metadatos["escribano"] == "Javier de Taboada Vizcarra"

    def test_carga_dos_registros_y_descarta_fantasmas(self):
        excel = ExcelRepository(self.repo.cargar("notarial"))
        fila = excel.detectar_fila_inicio_datos(self.tmp)
        records = excel.cargar_registros(self.tmp, fila, 0, 10 ** 6)

        assert len(records) == 2
        assert [r.fila for r in records] == [16, 17]
        assert records[0].escritura == "1"
        assert records[0].tomo == "1"
        assert records[0].titulo == "Protesto"
        assert records[0].fecha_inicio == "1967-12-20"
        assert records[0].fecha_descripcion == "2026-03-25"
        assert records[0].interesado2 == "Adelina Tejada Jara"
        assert records[1].folios == "1v - 2r"

    def test_guardar_registros_escribe_en_hoja_notarial(self):
        excel = ExcelRepository(self.repo.cargar("notarial"))
        fila = excel.detectar_fila_inicio_datos(self.tmp)
        records = excel.cargar_registros(self.tmp, fila, 0, 10 ** 6)

        records[0].titulo = "TESTAMENTO"
        records[0].tomo = "9"
        records[0].observaciones = "NOTA DE PRUEBA"
        celdas = excel.guardar_registros(self.tmp, fila, records)

        assert celdas > 0
        reloaded = excel.cargar_registros(self.tmp, fila, 0, 10 ** 6)
        assert reloaded[0].titulo == "TESTAMENTO"
        assert reloaded[0].tomo == "9"
        assert reloaded[0].observaciones == "NOTA DE PRUEBA"
        assert reloaded[0].escribano == "Javier de Taboada Vizcarra"

    def test_jerarquia_notarial_incluye_tomo_antes_de_registro(self):
        builder = HierarchyBuilder(self.repo.cargar("notarial"))
        record = InventoryRecord(
            id="#0001", fila=16, registro="1", escribano="Javier de Taboada Vizcarra",
            protocolo="1", folios="1r", pg_pdf="", titulo="Protesto",
            fecha_inicio="1967-12-20", interesado1="Oscar Gutierrez",
            interesado2="Adelina Tejada Jara", escritura="1", tomo="1",
        )

        ruta = builder.construir_ruta(
            record, "out", "1", "Javier de Taboada Vizcarra", ""
        )
        niveles = ruta[len("out"):].strip(os.sep).split(os.sep)

        assert len(niveles) == 12
        assert niveles[0] == "ACERVO DOCUMENTAL NUMERO 1"
        assert niveles[1] == "SIGLO 20"
        assert niveles[6] == "TOMO 1"
        assert niveles[7] == "REGISTRO 1"
        assert niveles[8] == "PROTESTO"
        assert niveles[9] == "12. DICIEMBRE"
        assert niveles[10] == "Oscar Gutierrez"
        assert niveles[11] == "Adelina Tejada Jara.pdf"

    def test_jerarquia_notarial_sin_tomo_usa_respaldo(self):
        builder = HierarchyBuilder(self.repo.cargar("notarial"))
        record = InventoryRecord(
            id="#0002", fila=17, registro="1", escribano="X",
            protocolo="1", folios="2r", pg_pdf="", titulo="Protesto",
            fecha_inicio="1967-12-20", escritura="2", tomo="",
        )
        ruta = builder.construir_ruta(record, "out", "1", "X", "")
        assert "SIN TOMO" in ruta
        assert "REGISTRO 2" in ruta
